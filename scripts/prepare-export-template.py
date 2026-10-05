"""Prepare a retained runtime DOCX from the supplied DOC's read-only conversion.
Usage: python scripts/prepare-export-template.py converted.docx output.docx
"""
import sys
from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document(sys.argv[1])
if "Title" not in doc.styles:
    title_style = doc.styles.add_style("Title", WD_STYLE_TYPE.PARAGRAPH)
    title_style.base_style = doc.styles["Normal"]
doc.paragraphs[0].style = doc.styles["Title"]
doc.core_properties.author = "WriteWise"
doc.core_properties.last_modified_by = "WriteWise"


def slot(cell, name):
    paragraph = cell.paragraphs[0]
    props = next((deepcopy(r._r.rPr) for r in paragraph.runs if r._r.rPr is not None), None)
    paragraph.clear()
    run = paragraph.add_run("{{" + name + "}}")
    if props is not None:
        run._r.insert(0, props)
    for extra in cell.paragraphs[1:]:
        cell._tc.remove(extra._p)


main = doc.tables[0]
support = doc.tables[1]
reflection = doc.tables[2]
slot(main.rows[1].cells[1], "TITLE")
slot(main.rows[1].cells[5], "DURATION")
slot(main.rows[2].cells[1], "GRADE_CLASS")
slot(main.rows[2].cells[3], "DESIGNER")
slot(main.rows[2].cells[5], "STUDENT_ID")
for row, token in [(4, "STANDARDS"), (5, "TEXTBOOK_ANALYSIS"), (6, "STUDENT_ANALYSIS"),
                   (8, "OBJECTIVES"), (9, "FOCUS"), (10, "DIFFICULTIES"),
                   (12, "DESIGN_RATIONALE"), (13, "FLOW")]:
    slot(main.rows[row].cells[-1], token)

# The source permits either a process table or prose. Keep its original table
# and remove instructional prose and the unused alternative example layout.
process_cell = main.rows[15].cells[0]
process_table = process_cell.tables[0]
for p in list(process_cell.paragraphs):
    process_cell._tc.remove(p._p)
process_cell.add_paragraph()
for cell, token in zip(process_table.rows[1].cells,
                       ["STAGE_NAME", "TEACHER_ACTIVITIES", "STUDENT_ACTIVITIES", "PURPOSE"]):
    slot(cell, token)
for row in process_table.rows:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
process_table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
header = process_table.rows[0].cells[0].paragraphs[0]
header.runs[0].text = "教学环节（用时）"
for run in header.runs[1:]:
    run.text = ""
# Word does not repeat headers for nested tables. Lift the original table to
# the document body so its four columns and repeated header work across pages.
main._tbl.addnext(process_table._tbl)
main._tbl.remove(main.rows[15]._tr)
process_table.alignment = WD_TABLE_ALIGNMENT.CENTER

slot(support.rows[2].cells[0], "BOARD_DESIGN")
support._tbl.remove(support.rows[1]._tr)
for row, token in [(3, "LEARNING_RESOURCES"), (4, "TEACHING_RESOURCES"),
                   (5, "ASSESSMENT"), (6, "HOMEWORK_REQUIREMENTS")]:
    slot(support.rows[row].cells[-1], token)
slot(reflection.rows[1].cells[0], "REFLECTION")

# Content controls row height; keep page geometry, grids, borders and fonts.
for height in list(doc._element.iter(qn("w:trHeight"))):
    if not any(height is h for h in reflection._tbl.iter(qn("w:trHeight"))):
        height.getparent().remove(height)
doc.save(sys.argv[2])
